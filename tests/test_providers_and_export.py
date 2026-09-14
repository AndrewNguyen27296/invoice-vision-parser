"""Provider wiring, the strategy ladder, and the Excel deliverable."""

from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from parser import extractor as EX
from parser.cost_guard import CostGuard, GuardConfig
from parser.providers import ProviderReply, available_providers, get_provider
from parser.providers.base import ProviderError
from parser.providers.mock_provider import MOCK_INVOICE
from utils.exporter import (
    build_batch_workbook,
    to_csv_text,
    to_xlsx_bytes,
    write_xlsx,
)

SAMPLES = Path(__file__).resolve().parent.parent / "sample_invoices"
DIGITAL = SAMPLES / "1_clean_freight.pdf"
SCAN = SAMPLES / "3_skewed_scan.pdf"


def guard(tmp_path: Path, **over) -> CostGuard:
    cfg = GuardConfig()
    cfg.cache_dir = tmp_path / ".cache"
    for k, v in over.items():
        setattr(cfg, k, v)
    return CostGuard(cfg)


# ------------------------------------------------------------------ registry
def test_every_provider_satisfies_the_contract():
    for name in available_providers():
        p = get_provider(name)
        assert p.name == name
        assert isinstance(p.free, bool)
        assert isinstance(p.describe(), str) and p.describe()
        assert callable(p.call)


def test_an_unknown_provider_names_the_valid_ones():
    with pytest.raises(ProviderError, match="Available"):
        get_provider("wishful-thinking")


def test_gemini_free_tier_flag_is_honest_about_the_data_terms(monkeypatch):
    monkeypatch.setenv("PARSER_GEMINI_FREE_TIER", "1")
    free = get_provider("gemini")
    assert free.free
    # The free tier's catch must be stated where an operator will read it.
    assert "improve their products" in free.describe()

    monkeypatch.setenv("PARSER_GEMINI_FREE_TIER", "0")
    paid = get_provider("gemini")
    assert not paid.free
    assert "not used for training" in paid.describe()


def test_gemini_without_a_key_fails_clearly_and_for_free(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    with pytest.raises(ProviderError, match="aistudio.google.com"):
        get_provider("gemini")._client()


# --------------------------------------------------------- strategy ladder
def test_a_digital_pdf_never_reaches_a_paid_model(tmp_path):
    """The whole economic argument, in one assertion."""
    called = {"n": 0}

    class Tripwire:
        name, free = "tripwire", False
        describe = staticmethod(lambda: "must not be called")

        def call(self, pages, cfg):
            called["n"] += 1
            raise AssertionError("a digital PDF must be read from its text layer")

    result = EX.extract_invoice(DIGITAL, guard=guard(tmp_path), provider=Tripwire())
    assert called["n"] == 0
    assert result.strategy == EX.STRATEGY_TEXT_LAYER
    assert result.cost_usd == 0.0 and result.was_free
    assert result.audit.passed
    assert "no API call" in result.route


def test_a_scan_does_reach_the_provider(tmp_path):
    class Recorder:
        name, free = "recorder", True
        describe = staticmethod(lambda: "recorder")

        def __init__(self):
            self.calls = 0

        def call(self, pages, cfg):
            self.calls += 1
            assert pages, "the provider must receive rendered pages"
            return ProviderReply(json.dumps(MOCK_INVOICE), 1200, 400, "recorder")

    rec = Recorder()
    result = EX.extract_invoice(SCAN, guard=guard(tmp_path), provider=rec)
    assert rec.calls == 1
    assert result.strategy == EX.STRATEGY_VISION
    assert result.provider == "recorder"


def test_disabling_the_text_layer_forces_vision(tmp_path):
    class Recorder:
        name, free = "rec", True
        describe = staticmethod(lambda: "rec")

        def call(self, pages, cfg):
            return ProviderReply(json.dumps(MOCK_INVOICE), 10, 5, "rec")

    result = EX.extract_invoice(DIGITAL, guard=guard(tmp_path),
                                provider=Recorder(), use_text_layer=False)
    assert result.strategy == EX.STRATEGY_VISION


def test_a_broken_text_layer_parser_falls_back_instead_of_crashing(tmp_path, monkeypatch):
    """A bug in the free path must cost accuracy, never availability."""
    monkeypatch.setattr(
        EX, "extract_from_text_layer",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("parser exploded")))

    class Fallback:
        name, free = "fallback", True
        describe = staticmethod(lambda: "fallback")

        def call(self, pages, cfg):
            return ProviderReply(json.dumps(MOCK_INVOICE), 1, 1, "fallback")

    result = EX.extract_invoice(DIGITAL, guard=guard(tmp_path), provider=Fallback())
    assert result.strategy == EX.STRATEGY_VISION
    assert any("text-layer read failed" in n for n in result.invoice.extraction_notes)


def test_free_routes_are_counted_but_not_charged(tmp_path):
    g = guard(tmp_path)
    EX.extract_invoice(DIGITAL, guard=g)
    today = g.today()
    assert today["pages"] == 1
    assert today["free_pages"] == 1
    assert today["cost_usd"] == 0.0
    assert "1 free" in g.summary()


# ------------------------------------------------------------------- export
def test_workbook_holds_real_numbers_a_formula_and_the_audit(tmp_path):
    result = EX.extract_invoice(DIGITAL, guard=guard(tmp_path))
    path = write_xlsx(result, tmp_path / "out.xlsx")
    ws = load_workbook(path).active

    text = "\n".join(
        str(ws.cell(row=r, column=c).value)
        for r in range(1, ws.max_row + 1) for c in range(1, 8)
        if ws.cell(row=r, column=c).value is not None
    )
    assert "NORDFRAKT LOGISTIK AB" in text
    assert "Math Audit Passed" in text
    assert text.count("MATH AUDIT") == 0, "the verdict must not be printed twice"
    assert "=SUM(" in text, "the sheet must re-add the column itself"

    amounts = [ws.cell(row=r, column=5).value for r in range(1, ws.max_row + 1)]
    numeric = [v for v in amounts if isinstance(v, (int, float))]
    assert 740.0 in numeric and 1937.5 in numeric
    for row in range(1, ws.max_row + 1):
        cell = ws.cell(row=row, column=5)
        if isinstance(cell.value, (int, float)):
            assert cell.number_format == "#,##0.00", "a clerk must be able to sum this"


def test_a_failed_audit_is_visible_in_the_export(tmp_path):
    result = EX.extract_invoice(SAMPLES / "2_messy_wholesale.pdf", guard=guard(tmp_path))
    ws = load_workbook(write_xlsx(result, tmp_path / "fail.xlsx")).active
    assert "FAILED" in str(ws["A2"].value)
    body = "\n".join(str(ws.cell(row=r, column=2).value) for r in range(1, ws.max_row + 1))
    assert "differs by 3.00" in body


def test_xlsx_bytes_are_a_real_workbook(tmp_path):
    result = EX.extract_invoice(DIGITAL, guard=guard(tmp_path))
    blob = to_xlsx_bytes(result)
    assert blob[:2] == b"PK" and len(blob) > 4000


def test_csv_export_is_clean(tmp_path):
    result = EX.extract_invoice(DIGITAL, guard=guard(tmp_path))
    csv_text = to_csv_text(result.invoice)
    lines = csv_text.strip().splitlines()
    assert lines[0] == "description,quantity,unit_price,line_total"
    assert len(lines) == 5
    assert '"Road freight Gothenburg -> Malmo, 4 pallets",4,185.00,740.00' in csv_text


def test_batch_sheet_is_one_row_per_document(tmp_path):
    g = guard(tmp_path)
    results = [
        EX.extract_invoice(SAMPLES / n, guard=g, mock=(n == "3_skewed_scan.pdf"))
        for n in ("1_clean_freight.pdf", "2_messy_wholesale.pdf", "3_skewed_scan.pdf")
    ]
    ws = build_batch_workbook(results).active
    assert ws.max_row == 4                      # header + 3
    verdicts = [ws.cell(row=r, column=9).value for r in range(2, 5)]
    assert verdicts[0] == "MATH VERIFIED"
    assert verdicts[1] == "MATH AUDIT FAILED"
    assert ws.cell(row=3, column=10).value == pytest.approx(3.00)
